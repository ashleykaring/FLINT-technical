from __future__ import annotations

import csv
import hashlib
import json
from collections.abc import Iterable, Sequence
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import duckdb
from openpyxl import load_workbook

from pipeline.ingest.contracts import SourceContract

METADATA_COLUMNS = (
    "_source_file",
    "_source_row",
    "_ingestion_run_id",
    "_ingested_at",
    "_record_hash",
)


class SourceContractError(ValueError):
    """Raised when an input does not match its declared source contract."""


def _quote_identifier(value: str) -> str:
    return f'"{value.replace(chr(34), chr(34) * 2)}"'


def _fingerprint(values: Iterable[Any]) -> str:
    """Create a duplicate-detection fingerprint; this is not a security hash."""

    content = "\x1f".join("" if value is None else str(value) for value in values)
    return hashlib.md5(content.encode("utf-8"), usedforsecurity=False).hexdigest()


def _as_raw_string(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(value)


def _read_csv_header(path: Path) -> tuple[str, ...]:
    with path.open("r", encoding="utf-8-sig", newline="") as source:
        return tuple(next(csv.reader(source)))


def _validate_columns(
    path: Path,
    actual: Sequence[str],
    expected: Sequence[str],
) -> None:
    if tuple(actual) != tuple(expected):
        raise SourceContractError(
            f"{path.name} columns do not match the source contract: "
            f"expected {list(expected)!r}, received {list(actual)!r}"
        )


def load_csv_source(
    connection: duckdb.DuckDBPyConnection,
    path: Path,
    contract: SourceContract,
    run_id: str,
) -> int:
    """Load a CSV with all source columns retained as VARCHAR values."""

    _validate_columns(path, _read_csv_header(path), contract.columns)
    table = _quote_identifier(contract.table_name)
    hash_values = ", ".join(
        f"coalesce({_quote_identifier(column)}, '')" for column in contract.columns
    )

    connection.execute(
        f"""
        CREATE OR REPLACE TABLE {table} AS
        SELECT
            *,
            ?::VARCHAR AS _source_file,
            (row_number() OVER () + 1)::BIGINT AS _source_row,
            ?::VARCHAR AS _ingestion_run_id,
            current_timestamp AS _ingested_at,
            md5(concat_ws(chr(31), {hash_values})) AS _record_hash
        FROM read_csv(?, header = true, all_varchar = true)
        """,
        [path.name, run_id, str(path)],
    )
    return _table_count(connection, contract.table_name)


def _create_raw_table(
    connection: duckdb.DuckDBPyConnection,
    table_name: str,
    business_columns: Sequence[str],
) -> None:
    definitions = [f"{_quote_identifier(column)} VARCHAR" for column in business_columns]
    definitions.extend(
        [
            "_source_file VARCHAR",
            "_source_row BIGINT",
            "_ingestion_run_id VARCHAR",
            "_ingested_at TIMESTAMPTZ",
            "_record_hash VARCHAR",
        ]
    )
    connection.execute(
        f"CREATE OR REPLACE TABLE {_quote_identifier(table_name)} ({', '.join(definitions)})"
    )


def _insert_rows(
    connection: duckdb.DuckDBPyConnection,
    table_name: str,
    rows: Sequence[Sequence[Any]],
) -> None:
    if not rows:
        return
    placeholders = ", ".join("?" for _ in rows[0])
    connection.executemany(
        f"INSERT INTO {_quote_identifier(table_name)} VALUES ({placeholders})",
        rows,
    )


def load_estimates_api(
    connection: duckdb.DuckDBPyConnection,
    path: Path,
    run_id: str,
) -> tuple[int, int]:
    with path.open("r", encoding="utf-8") as source:
        payload = json.load(source)

    if not isinstance(payload, dict) or not isinstance(payload.get("data"), list):
        raise SourceContractError(f"{path.name} must contain an object with a data array")

    projects = payload["data"]
    declared_count = payload.get("count")
    if declared_count != len(projects):
        raise SourceContractError(
            f"{path.name} declares {declared_count!r} projects but contains {len(projects)}"
        )

    project_columns = ("project_id", "line_item_count", "raw_json")
    line_columns = (
        "project_id",
        "cost_code",
        "description",
        "original_estimate",
        "current_estimate",
        "raw_json",
    )
    _create_raw_table(connection, "raw_estimate_api_projects", project_columns)
    _create_raw_table(connection, "raw_estimate_api_line_items", line_columns)

    loaded_at = datetime.now(UTC)
    project_rows: list[tuple[Any, ...]] = []
    line_rows: list[tuple[Any, ...]] = []
    line_number = 0

    for project_number, project in enumerate(projects, start=1):
        if not isinstance(project, dict) or not isinstance(project.get("line_items"), list):
            raise SourceContractError(
                f"{path.name} project {project_number} is missing a line_items array"
            )

        project_id = _as_raw_string(project.get("project_id"))
        line_items = project["line_items"]
        project_values = (
            project_id,
            str(len(line_items)),
            json.dumps(project, ensure_ascii=False, sort_keys=True),
        )
        project_rows.append(
            (
                *project_values,
                path.name,
                project_number,
                run_id,
                loaded_at,
                _fingerprint(project_values),
            )
        )

        for item in line_items:
            if not isinstance(item, dict):
                raise SourceContractError(
                    f"{path.name} project {project_number} contains a non-object line item"
                )
            line_number += 1
            line_values = (
                project_id,
                _as_raw_string(item.get("cost_code")),
                _as_raw_string(item.get("description")),
                _as_raw_string(item.get("original_estimate")),
                _as_raw_string(item.get("current_estimate")),
                json.dumps(item, ensure_ascii=False, sort_keys=True),
            )
            line_rows.append(
                (
                    *line_values,
                    path.name,
                    line_number,
                    run_id,
                    loaded_at,
                    _fingerprint(line_values[:-1]),
                )
            )

    _insert_rows(connection, "raw_estimate_api_projects", project_rows)
    _insert_rows(connection, "raw_estimate_api_line_items", line_rows)
    return len(project_rows), len(line_rows)


def load_equipment(
    connection: duckdb.DuckDBPyConnection,
    path: Path,
    contract: SourceContract,
    run_id: str,
) -> int:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        source_rows = sheet.iter_rows(
            min_row=5,
            min_col=2,
            max_col=8,
            values_only=True,
        )
        actual_headers = tuple(next(source_rows))
        _validate_columns(path, actual_headers, contract.columns)
        _create_raw_table(connection, contract.table_name, contract.columns)

        loaded_at = datetime.now(UTC)
        rows: list[tuple[Any, ...]] = []
        for source_row, raw_values in enumerate(source_rows, start=6):
            values = tuple(_as_raw_string(value) for value in raw_values)
            if all(value is None or not value.strip() for value in values):
                continue
            rows.append(
                (
                    *values,
                    path.name,
                    source_row,
                    run_id,
                    loaded_at,
                    _fingerprint(values),
                )
            )

        _insert_rows(connection, contract.table_name, rows)
        return len(rows)
    finally:
        workbook.close()


def load_correspondence(
    connection: duckdb.DuckDBPyConnection,
    path: Path,
    contract: SourceContract,
    run_id: str,
) -> int:
    business_columns = (*contract.columns, "raw_json")
    _create_raw_table(connection, contract.table_name, business_columns)
    loaded_at = datetime.now(UTC)
    rows: list[tuple[Any, ...]] = []

    with path.open("r", encoding="utf-8") as source:
        for source_row, raw_line in enumerate(source, start=1):
            if not raw_line.strip():
                continue
            try:
                item = json.loads(raw_line)
            except json.JSONDecodeError as exc:
                raise SourceContractError(
                    f"{path.name} line {source_row} is not valid JSON: {exc.msg}"
                ) from exc

            if not isinstance(item, dict):
                raise SourceContractError(f"{path.name} line {source_row} is not a JSON object")
            missing = [column for column in contract.columns if column not in item]
            if missing:
                raise SourceContractError(
                    f"{path.name} line {source_row} is missing columns {missing!r}"
                )

            values = tuple(_as_raw_string(item.get(column)) for column in contract.columns)
            rows.append(
                (
                    *values,
                    raw_line.rstrip("\r\n"),
                    path.name,
                    source_row,
                    run_id,
                    loaded_at,
                    _fingerprint(values),
                )
            )

    _insert_rows(connection, contract.table_name, rows)
    return len(rows)


def validate_estimate_equivalence(connection: duckdb.DuckDBPyConnection) -> None:
    differences = connection.execute(
        """
        WITH api AS (
            SELECT
                project_id,
                cost_code,
                description,
                original_estimate,
                coalesce(current_estimate, '') AS current_estimate
            FROM raw_estimate_api_line_items
        ),
        csv AS (
            SELECT
                project AS project_id,
                est_cost_code AS cost_code,
                description,
                original_estimate,
                coalesce(current_estimate, '') AS current_estimate
            FROM raw_estimate_csv
        ),
        differences AS (
            (SELECT * FROM api EXCEPT SELECT * FROM csv)
            UNION ALL
            (SELECT * FROM csv EXCEPT SELECT * FROM api)
        )
        SELECT count(*) FROM differences
        """
    ).fetchone()[0]
    if differences:
        raise SourceContractError(
            "estimates_budget.csv does not match estimates_budget_api.json "
            f"({differences} differing records)"
        )


def _table_count(connection: duckdb.DuckDBPyConnection, table_name: str) -> int:
    return connection.execute(f"SELECT count(*) FROM {_quote_identifier(table_name)}").fetchone()[0]
